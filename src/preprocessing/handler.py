"""Preprocessing Lambda: convert heterogeneous uploads to normalized text.

Triggered by S3 events on the raw uploads bucket or by direct invocation.
Detects content type and runs the appropriate converter:
  - PDF -> PyPDF2 text extraction
  - Image -> Amazon Nova Lite multimodal description
  - URL -> httpx fetch + trafilatura extraction
  - DOCX -> python-docx paragraph extraction
  - Text -> UTF-8 passthrough

Output: normalized text written to KB source bucket under tenant prefix.
"""

from __future__ import annotations

import json
import logging
import os
import re
import uuid
from typing import Any

import boto3

from common import config
from common.logging_config import setup_logging

setup_logging()
logger = logging.getLogger(__name__)

_s3 = None
_bedrock = None


def _get_s3():
    global _s3
    if _s3 is None:
        _s3 = boto3.client("s3", region_name=config.AWS_REGION)
    return _s3


def _get_bedrock():
    global _bedrock
    if _bedrock is None:
        _bedrock = boto3.client("bedrock-runtime", region_name=config.AWS_REGION)
    return _bedrock


CONTENT_TYPE_MAP = {
    ".pdf": "pdf",
    ".png": "image",
    ".jpg": "image",
    ".jpeg": "image",
    ".gif": "image",
    ".webp": "image",
    ".docx": "docx",
    ".txt": "text",
    ".md": "text",
    ".csv": "text",
    ".json": "text",
}


def handler(event: dict[str, Any], context: Any = None) -> dict[str, Any]:
    if "Records" in event:
        return _handle_s3_event(event)
    return _handle_direct(event)


def _handle_s3_event(event: dict) -> dict:
    results = []
    for record in event.get("Records", []):
        bucket = record["s3"]["bucket"]["name"]
        key = record["s3"]["object"]["key"]
        try:
            result = _process_object(bucket, key)
            results.append(result)
        except Exception:
            logger.exception("Failed to process s3://%s/%s", bucket, key)
            results.append({"key": key, "status": "error"})
    return {"results": results}


def _handle_direct(event: dict) -> dict:
    tenant_id = event.get("tenant_id", "")
    source_type = event.get("source_type", "text")
    content = event.get("content", "")
    url = event.get("url", "")
    doc_name = event.get("doc_name", f"note_{uuid.uuid4().hex[:8]}")
    raw_bucket = event.get("bucket", config.RAW_UPLOADS_BUCKET)
    raw_key = event.get("key", "")

    if not tenant_id:
        return {"error": "tenant_id required"}

    if source_type == "url":
        if not url:
            return {"error": "url required for url source_type"}
        text = _extract_url(url)
        method = "trafilatura"
    elif raw_key:
        raw_bytes = _get_s3().get_object(Bucket=raw_bucket, Key=raw_key)["Body"].read()
        ext = os.path.splitext(raw_key)[1].lower()
        detected_type = CONTENT_TYPE_MAP.get(ext, "text")
        text, method = _convert(raw_bytes, detected_type, ext)
    elif content:
        text = content
        method = "passthrough"
    else:
        return {"error": "content, url, or key required"}

    if not text or not text.strip():
        return {"error": "no text extracted", "method": method}

    normalized_key = f"{tenant_id}/{doc_name}.txt"
    _get_s3().put_object(
        Bucket=config.KB_SOURCE_BUCKET,
        Key=normalized_key,
        Body=text.encode("utf-8"),
        ContentType="text/plain",
        Metadata={
            "tenant_id": tenant_id,
            "source_type": source_type,
            "extraction_method": method,
        },
    )

    _trigger_kb_reingestion()

    return {
        "status": "ok",
        "tenant_id": tenant_id,
        "normalized_key": normalized_key,
        "char_count": len(text),
        "extraction_method": method,
    }


def _process_object(bucket: str, key: str) -> dict:
    parts = key.split("/", 1)
    if len(parts) < 2:
        return {"key": key, "status": "skipped", "reason": "no tenant prefix"}

    tenant_id = parts[0]
    filename = parts[-1]
    ext = os.path.splitext(filename)[1].lower()
    detected_type = CONTENT_TYPE_MAP.get(ext, "text")

    raw_bytes = _get_s3().get_object(Bucket=bucket, Key=key)["Body"].read()
    text, method = _convert(raw_bytes, detected_type, ext)

    if not text or not text.strip():
        logger.warning("No text extracted from s3://%s/%s", bucket, key)
        return {"key": key, "status": "empty", "method": method}

    doc_name = os.path.splitext(filename)[0]
    normalized_key = f"{tenant_id}/{doc_name}.txt"
    _get_s3().put_object(
        Bucket=config.KB_SOURCE_BUCKET,
        Key=normalized_key,
        Body=text.encode("utf-8"),
        ContentType="text/plain",
        Metadata={
            "tenant_id": tenant_id,
            "source_type": detected_type,
            "extraction_method": method,
            "original_key": key,
        },
    )
    _trigger_kb_reingestion()
    return {"key": key, "status": "ok", "normalized_key": normalized_key, "char_count": len(text), "method": method}


def _convert(raw_bytes: bytes, source_type: str, ext: str) -> tuple[str, str]:
    if source_type == "pdf":
        return _extract_pdf(raw_bytes), "pypdf2"
    elif source_type == "image":
        return _extract_image(raw_bytes, ext), "nova_lite"
    elif source_type == "docx":
        return _extract_docx(raw_bytes), "python_docx"
    else:
        return raw_bytes.decode("utf-8", errors="replace"), "passthrough"


def _extract_pdf(data: bytes) -> str:
    try:
        import PyPDF2
        import io
        reader = PyPDF2.PdfReader(io.BytesIO(data))
        pages = []
        for page in reader.pages:
            text = page.extract_text()
            if text:
                pages.append(text)
        return "\n\n".join(pages)
    except Exception:
        logger.exception("PyPDF2 extraction failed")
        return data.decode("utf-8", errors="replace")


def _extract_image(data: bytes, ext: str) -> str:
    import base64
    media_type_map = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".gif": "image/gif", ".webp": "image/webp"}
    b64_data = base64.b64encode(data).decode("utf-8")
    body = json.dumps({
        "messages": [{
            "role": "user",
            "content": [
                {"image": {"format": ext.lstrip("."), "source": {"bytes": b64_data}}},
                {"text": "Extract ALL text and information from this image. If it contains text, transcribe it exactly. If it contains visual content, describe it in detail. Output only the extracted content."},
            ],
        }],
        "inferenceConfig": {"maxTokens": 2000},
    })
    try:
        response = _get_bedrock().invoke_model(
            modelId=config.NOVA_MODEL_ID, contentType="application/json", accept="application/json", body=body.encode("utf-8"),
        )
        result = json.loads(response["body"].read())
        content_blocks = result.get("output", {}).get("message", {}).get("content", [])
        return "\n".join(b["text"] for b in content_blocks if "text" in b)
    except Exception:
        logger.exception("Nova Lite image extraction failed")
        return ""


def _extract_url(url: str) -> str:
    html = ""
    try:
        import httpx
        resp = httpx.get(url, timeout=30, follow_redirects=True)
        resp.raise_for_status()
        html = resp.text
    except ImportError:
        logger.info("httpx not available, using urllib")
        import urllib.request
        req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0 LinkMe-Bot/1.0"})
        try:
            with urllib.request.urlopen(req, timeout=30) as resp:
                html = resp.read().decode("utf-8", errors="replace")
        except Exception:
            logger.exception("urllib fetch failed for %s", url)
            return ""
    except Exception:
        logger.exception("Failed to fetch URL: %s", url)
        return ""

    if not html:
        return ""

    try:
        import trafilatura
        text = trafilatura.extract(html)
        if text:
            return text
    except ImportError:
        logger.info("trafilatura not available, using basic HTML strip")
    except Exception:
        logger.exception("trafilatura extraction failed for %s", url)
    return _basic_html_strip(html)


def _basic_html_strip(html: str) -> str:
    text = re.sub(r"<script[^>]*>.*?</script>", "", html, flags=re.DOTALL | re.IGNORECASE)
    text = re.sub(r"<style[^>]*>.*?</style>", "", text, flags=re.DOTALL | re.IGNORECASE)
    text = re.sub(r"<[^>]+>", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text


def _extract_docx(data: bytes) -> str:
    try:
        import docx
        import io
        doc = docx.Document(io.BytesIO(data))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        tables_text = []
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    tables_text.append(" | ".join(cells))
        return "\n\n".join(paragraphs + tables_text)
    except Exception:
        logger.exception("python-docx extraction failed")
        return data.decode("utf-8", errors="replace")


def _trigger_kb_reingestion() -> str | None:
    try:
        bedrock_agent = boto3.client("bedrock-agent", region_name=config.AWS_REGION)
        kb_id = config.BEDROCK_KB_ID
        if not kb_id:
            return None
        ds_id = os.environ.get("KB_DS_ID", "EARAY9GAXN")
        resp = bedrock_agent.start_ingestion_job(knowledgeBaseId=kb_id, dataSourceId=ds_id)
        job_id = resp.get("ingestionJob", {}).get("ingestionJobId", "unknown")
        logger.info("KB re-ingestion triggered: job=%s", job_id)
        return job_id
    except Exception:
        logger.exception("Failed to trigger KB re-ingestion")
        return None
